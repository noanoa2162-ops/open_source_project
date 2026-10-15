"""Excel export — one row per commitment form.

Everything is written as text. Identifiers keep their leading zeros, dates are
not re-parsed into Excel serial numbers, and a value starting with a formula
character can never execute when the lab opens the file.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.fields import FIELDS, MULTI_SEPARATOR
from app.models import STATE_REVIEW, Batch, FormRecord
from app.validators import REASON_TEXT

# נוסחאות: ערך שמתחיל באחד מאלה ייכתב כטקסט מנוטרל ולא ירוץ כנוסחה.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

NAVY = "1B3A5C"
AMBER_FILL = "FFF1DC"
AMBER_TEXT = "8A4B10"
LINE = "D8E0E8"

EXTRA_COLUMNS = ("קובץ מקור", "עמוד בקובץ", "שדות לבדיקה", "סיבות")


def neutralize(value: str) -> str:
    """Stop a spreadsheet from treating a scanned value as a formula."""
    if value and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def record_row(record: FormRecord) -> list[str]:
    cells = [MULTI_SEPARATOR.join(record.fields[f.key].values) for f in FIELDS]
    flagged = [f for f in record.fields.values() if f.state == STATE_REVIEW]
    reasons = sorted({REASON_TEXT.get(r, r) for f in flagged for r in f.reasons})
    # שדה במחלוקת נשאר ריק בתא — אף קורא לא מיוצא כעובדה. האפשרויות מופיעות כאן.
    for f in flagged:
        if f.candidates:
            label = next(s.label_he for s in FIELDS if s.key == f.key)
            reasons.append(f"{label}: " + " / ".join(f.candidates))
    return [
        *cells,
        record.source_name,
        str(record.page_number),
        str(len(flagged)),
        "; ".join(reasons),
    ]


def build_workbook(batch: Batch) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "התחייבויות"
    ws.sheet_view.rightToLeft = True

    headers = [f.label_he for f in FIELDS] + list(EXTRA_COLUMNS)
    widths = [14, 16, 16, 18, 20, 22, 20, 30, 12, 12, 46]
    thin = Side(style="thin", color=LINE)

    for index, (label, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=index, value=label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
        ws.column_dimensions[get_column_letter(index)].number_format = "@"
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for offset, record in enumerate(batch.records):
        row_index = offset + 2
        for column_index, value in enumerate(record_row(record), start=1):
            cell = ws.cell(row=row_index, column=column_index, value=neutralize(value))
            cell.number_format = "@"  # טקסט: אפסים מובילים ותאריכים נשמרים כפי שהם
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)

            spec = FIELDS[column_index - 1] if column_index <= len(FIELDS) else None
            if spec and record.fields[spec.key].state == STATE_REVIEW:
                # כתום = נדרשת בדיקה. אותה משמעות כמו במסך.
                cell.fill = PatternFill("solid", fgColor=AMBER_FILL)
                cell.font = Font(color=AMBER_TEXT)

    _add_notice(ws, len(batch.records) + 3, batch)
    return wb


def _add_notice(ws, row: int, batch: Batch) -> None:
    sources = len(batch.accepted_sources)
    lines = [
        "דמו להאקתון — מיועד לטפסי הדוגמה בלבד. אין להסתמך על הקובץ ללא בדיקה.",
        f"מקור: {batch.display_name} · {sources} קבצים · {len(batch.records)} טפסים · "
        f"{batch.review_total} שדות מסומנים לבדיקה",
        "תאים כתומים דורשים בדיקה אנושית, כולל שדות ריקים. אין כאן טענת נכונות.",
    ]
    if batch.rejected_sources:
        names = ", ".join(f"{s.filename} ({s.error})" for s in batch.rejected_sources)
        lines.append(f"קבצים שלא עובדו: {names}")
    for offset, text in enumerate(lines):
        cell = ws.cell(row=row + offset, column=1, value=text)
        cell.font = Font(size=10, color="5A6B7B", italic=offset > 0)
        cell.alignment = Alignment(horizontal="right")


def workbook_bytes(batch: Batch) -> bytes:
    buffer = BytesIO()
    build_workbook(batch).save(buffer)
    return buffer.getvalue()
