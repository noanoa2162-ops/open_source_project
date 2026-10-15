"""Acceptance checks for version 2: two judges, the secretary decides.

Every test here maps to a line in SPECIFICATION.md §14. No network: judges are
stand-ins that return exactly what each test needs.
"""

from __future__ import annotations

from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import gold, pipeline, store
from app.consensus import combine, normalize
from app.crops import crop_field, padded
from app.entities import canonical_sender
from app.excel import workbook_bytes
from app.fields import FIELDS
from app.models import STATE_APPROVED, STATE_OK, STATE_REVIEW, Batch, FormRecord, SourceFile
from app.pipeline import Panel, read_form
from app.providers.base import ExtractionError, RawExtraction, clamp_region
from app.validators import validate_record

SOURCE = SourceFile(index=1, filename="form.pdf", kind="pdf", sha256="")
ALL_CLEAR = {f.key: "clear" for f in FIELDS}


def reading(model="a", legibility=None, regions=None, context=None, **values) -> RawExtraction:
    return RawExtraction(
        values={f.key: values.get(f.key, []) for f in FIELDS},
        model=model,
        legibility={**ALL_CLEAR, **(legibility or {})},
        regions=regions or {},
        context=context or {},
    )


def record_from(*readings: RawExtraction, low_res=False, today=date(2026, 10, 14)) -> FormRecord:
    record = FormRecord.empty("b", SOURCE, 1)
    record.low_resolution = low_res
    combine(record, list(readings))
    validate_record(record, today=today)
    return record


class Judge:
    """A stand-in judge: fixed answer, counts its calls, can fail on demand."""

    def __init__(self, model, answer=None, error=None):
        self.name = self.model = model
        self.answer = answer
        self.error = error
        self.calls = 0

    def extract(self, image_path):
        self.calls += 1
        if self.error:
            raise self.error
        return self.answer


def png(path: Path, size=(1700, 2400)) -> Path:
    Image.new("RGB", size, "white").save(path, format="PNG")
    return path


# ------------------------------------------------------------- agreement


def test_agreement_gives_a_neutral_value_never_green():
    r = record_from(
        reading("a", commitment_numbers=["730578626"]),
        reading("b", commitment_numbers=["730578626"]),
    )
    field = r.fields["commitment_numbers"]
    assert field.values == ["730578626"]
    assert field.state == STATE_OK, "הסכמה = ניטרלי"
    assert field.state != STATE_APPROVED, "ירוק רק באישור אנושי"


def test_disagreement_is_orange_with_both_options_and_no_value():
    r = record_from(
        reading("a", national_id=["001122344"]),
        reading("b", national_id=["00112?344"]),
    )
    field = r.fields["national_id"]
    assert field.values == [], "אף קורא לא מוצג כעובדה"
    assert field.candidates == ["001122344", "00112?344"]
    assert field.state == STATE_REVIEW
    assert "readers_disagree" in field.reasons
    assert "missing" not in field.reasons, "'ריק' היה מטעה כשיש אפשרויות"


def test_only_one_judge_finding_a_value_is_orange_with_that_option():
    r = record_from(reading("a", first_name=["שמואל"]), reading("b"))
    field = r.fields["first_name"]
    assert field.values == []
    assert field.candidates == ["שמואל"]
    assert "single_value" in field.reasons


def test_both_empty_is_orange_missing():
    r = record_from(reading("a"), reading("b"))
    assert "missing" in r.fields["last_name"].reasons


def test_question_mark_is_orange_even_when_both_agree():
    r = record_from(
        reading("a", national_id=["0012?3456"]),
        reading("b", national_id=["0012?3456"]),
    )
    field = r.fields["national_id"]
    assert field.values == ["0012?3456"], "ה-? נשאר גלוי"
    assert "illegible" in field.reasons


def test_a_judge_admitting_doubt_makes_agreement_orange():
    r = record_from(
        reading("a", legibility={"first_name": "partial"}, first_name=["שמואל"]),
        reading("b", first_name=["שמואל"]),
    )
    assert r.fields["first_name"].state == STATE_REVIEW
    assert "illegible" in r.fields["first_name"].reasons


def test_one_reader_marks_every_field():
    r = record_from(reading("a", commitment_numbers=["730578626"]))
    assert all("single_reader" in f.reasons for f in r.fields.values())


def test_multi_values_agree_regardless_of_order_and_keep_first_order():
    r = record_from(
        reading("a", commitment_numbers=["111", "222"]),
        reading("b", commitment_numbers=["222", "111"]),
    )
    assert r.fields["commitment_numbers"].values == ["111", "222"]


# ---------------------------------------------------------- normalization


@pytest.mark.parametrize(
    ("key", "a", "b"),
    [
        ("sample_received_date", "18/06/2026", "18-06-2026"),
        ("sample_received_date", "18.6.26", "18/06/2026"),
        ("national_id", "001 122 334", "001-122-334"),
        ("sending_entity", "קופת חולים מאוחדת", "מאוחדת"),
        ("sending_entity", "המרכז הרפואי שמיר (אסף הרופא)", "אסף הרופא"),
        ("first_name", "שָׁלוֹם", "שלום"),
    ],
)
def test_differences_that_are_not_differences_still_agree(key, a, b):
    assert normalize(key, a) == normalize(key, b)


def test_normalization_never_changes_a_digit():
    assert normalize("national_id", "0012") == "0012"
    assert normalize("national_id", "001122344") != normalize("national_id", "001122334")


def test_sender_is_shown_by_its_canonical_name():
    r = record_from(
        reading("a", sending_entity=["קופת חולים מאוחדת"]),
        reading("b", sending_entity=["מאוחדת"]),
    )
    assert r.fields["sending_entity"].values == ["מאוחדת"]
    assert r.fields["sending_entity"].state == STATE_OK


def test_patient_sticker_kupah_is_not_confused_with_sender_list_lookup():
    # "מכבי" is a known kupah; if both judges wrongly took it, the list can't save us —
    # but an unknown string is always flagged, never guessed into the list.
    assert canonical_sender("מכבי") == "מכבי"
    assert canonical_sender("מעבדה פרטית כלשהי") is None


@pytest.mark.parametrize("text", ["מרפאת זיו", "ד\"ר מאיר כהן", "קליניקת כרמל"])
def test_ordinary_words_do_not_become_hospitals(text):
    assert canonical_sender(text) is None


@pytest.mark.parametrize(
    ("text", "expected"),
    [("מאוחדת רמה אחרת", "מאוחדת"), ("המרכז הרפואי שמיר (אסף הרופא) מוסמך", "שמיר (אסף הרופא)"),
     ("בית החולים הציבורי אסותא אשדוד", "אסותא אשדוד"), ("זיו", "זיו")],
)
def test_known_senders_inside_longer_printed_names(text, expected):
    assert canonical_sender(text) == expected


def test_unknown_sender_is_kept_and_flagged():
    r = record_from(
        reading("a", sending_entity=["מרפאת הדר"]), reading("b", sending_entity=["מרפאת הדר"])
    )
    field = r.fields["sending_entity"]
    assert field.values == ["מרפאת הדר"]
    assert "unknown_entity" in field.reasons


# ------------------------------------------------------------- date rules


def test_received_before_issue_is_flagged():
    ctx = {"issue_date": "15/06/2026"}
    r = record_from(
        reading("a", context=ctx, sample_received_date=["10/06/2026"]),
        reading("b", context=ctx, sample_received_date=["10/06/2026"]),
    )
    assert "date_before_issue" in r.fields["sample_received_date"].reasons


def test_received_in_the_future_is_flagged():
    r = record_from(
        reading("a", sample_received_date=["01/01/2030"]),
        reading("b", sample_received_date=["01/01/2030"]),
    )
    assert "date_in_future" in r.fields["sample_received_date"].reasons


def test_received_equal_to_birth_date_is_flagged():
    ctx = {"birth_date": "04/12/2009"}
    r = record_from(
        reading("a", context=ctx, sample_received_date=["04/12/2009"]),
        reading("b", context=ctx, sample_received_date=["04/12/2009"]),
    )
    assert "date_is_birth" in r.fields["sample_received_date"].reasons


def test_context_date_is_ignored_when_judges_disagree_on_it():
    r = record_from(
        reading("a", context={"issue_date": "15/06/2026"}, sample_received_date=["10/06/2026"]),
        reading("b", context={"issue_date": "05/06/2026"}, sample_received_date=["10/06/2026"]),
    )
    assert "issue_date" not in r.context
    assert "date_before_issue" not in r.fields["sample_received_date"].reasons


def test_valid_received_date_passes():
    ctx = {"issue_date": "15/06/2026", "birth_date": "04/12/2009"}
    r = record_from(
        reading("a", context=ctx, sample_received_date=["18/06/2026"]),
        reading("b", context=ctx, sample_received_date=["18/06/2026"]),
    )
    assert r.fields["sample_received_date"].state == STATE_OK


# ------------------------------------------------------ input quality gate


def test_low_resolution_flags_handwriting_fields_even_when_judges_agree():
    r = record_from(
        reading("a", first_name=["שרה"], commitment_numbers=["730578626"]),
        reading("b", first_name=["שרה"], commitment_numbers=["730578626"]),
        low_res=True,
    )
    assert "low_resolution" in r.fields["first_name"].reasons
    assert "low_resolution" not in r.fields["commitment_numbers"].reasons


def test_screenshot_sized_image_is_detected_as_low_resolution(tmp_path):
    small = png(tmp_path / "shot.png", size=(631, 896))
    record = read_form(Panel([Judge("a", reading("a"))]), small, "b", SOURCE, 1)
    assert record.low_resolution


# ------------------------------------------------------------ the panel


def test_both_judges_are_asked_and_compared(tmp_path):
    a = Judge("a", reading("a", commitment_numbers=["730578626"]))
    b = Judge("b", reading("b", commitment_numbers=["730578626"]))
    record = read_form(Panel([a, b]), png(tmp_path / "p.png"), "b", SOURCE, 1)
    assert (a.calls, b.calls) == (1, 1)
    assert record.readers == ["a", "b"]
    assert record.fields["commitment_numbers"].values == ["730578626"]


def test_one_judge_failing_falls_back_to_the_other_visibly(tmp_path):
    a = Judge("a", error=ExtractionError("timeout"))
    b = Judge("b", reading("b", commitment_numbers=["730578626"]))
    record = read_form(Panel([a, b]), png(tmp_path / "p.png"), "b", SOURCE, 1)
    assert record.error == "", "הטופס לא נופל"
    assert all("single_reader" in f.reasons for f in record.fields.values())
    assert "timeout" in record.model_note


def test_account_failure_retires_the_judge_for_the_rest_of_the_batch(tmp_path):
    a = Judge("a", error=ExtractionError("אין קרדיט", fatal=True))
    b = Judge("b", reading("b"))
    panel = Panel([a, b])
    image = png(tmp_path / "p.png")
    read_form(panel, image, "b", SOURCE, 1)
    read_form(panel, image, "b", SOURCE, 2)
    assert a.calls == 1, "לא שורפים עוד קריאות על חשבון בלי קרדיט"
    assert b.calls == 2


def test_same_file_twice_costs_nothing_the_second_time(tmp_path):
    pipeline._READ_CACHE.clear()
    source = SourceFile(index=1, filename="f.pdf", sha256="abc123")
    answer = reading("a", commitment_numbers=["1"])
    answer.usage = __import__("app.providers.base", fromlist=["Usage"]).Usage(100, 10, 0.01)
    a = Judge("a", answer)
    image = png(tmp_path / "p.png")
    first = read_form(Panel([a]), image, "b", source, 1)
    second = read_form(Panel([a]), image, "b", source, 1)
    assert a.calls == 1
    assert first.cost_usd == pytest.approx(0.01)
    assert second.cost_usd == 0.0
    pipeline._READ_CACHE.clear()


def test_budget_cap_stops_before_the_next_call(tmp_path, monkeypatch):
    from app.providers.base import Usage

    monkeypatch.setattr("app.config.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.store.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr(
        pipeline, "SETTINGS",
        replace(pipeline.SETTINGS, max_cost_per_batch_usd=0.05, max_parallel_calls=1),
    )
    expensive = reading("a")
    expensive.usage = Usage(1, 1, 0.03)
    judge = Judge("a", expensive)
    monkeypatch.setattr(pipeline, "build_judges", lambda: ([judge], []))

    batch = Batch(batch_id="bat_budget")
    for i in range(1, 6):
        batch.sources.append(SourceFile(index=i, filename=f"{i}.png", kind="png"))
        (store.source_dir("bat_budget", i) / "source").write_bytes(
            png(tmp_path / f"{i}.png", size=(300, 400)).read_bytes()
        )
    store.put(batch)
    pipeline._process("bat_budget")

    assert judge.calls == 2, "0.03 + 0.03 עובר את 0.05 — הקריאה השלישית לא נשלחת"
    assert batch.budget_skipped_pages == 3
    assert batch.status == "partial"
    assert any("תקרת התקציב" in w for w in batch.warnings)


# ---------------------------------------------------------- regions & crops


@pytest.mark.parametrize("bad", [None, [1, 2], ["x", 0, 1, 1], [0.5, 0.5, 0.4, 0.6], [0, 0, 5, 5]])
def test_garbage_coordinates_are_dropped_not_repaired(bad):
    assert clamp_region(bad) is None


def test_crop_margins_stay_inside_the_page():
    for region in [(0.0, 0.0, 0.05, 0.02), (0.95, 0.97, 1.0, 1.0), (0.4, 0.4, 0.6, 0.45)]:
        x0, y0, x1, y1 = padded(region)
        assert 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1
        assert x0 <= region[0] and x1 >= region[2], "השוליים מכילים את השדה"


def test_crop_is_enlarged_and_cached(tmp_path):
    page = png(tmp_path / "page.png", size=(1000, 1400))
    target = tmp_path / "crops" / "001-national_id.png"
    out = crop_field(page, (0.4, 0.4, 0.6, 0.45), target)
    with Image.open(out) as im:
        assert im.width > 0.2 * 1000, "מוגדל, לא רק נחתך"
    stamp = out.stat().st_mtime_ns
    crop_field(page, (0.4, 0.4, 0.6, 0.45), target)
    assert out.stat().st_mtime_ns == stamp


def test_quarters_cover_the_whole_page_with_overlap():
    from app.providers.openai_vision import TILE_ORIGINS, TILE_SPAN

    assert TILE_SPAN > 0.5, "חפיפה במרכז"
    for ox, oy in TILE_ORIGINS:
        assert 0 <= ox and ox + TILE_SPAN <= 1.0001
        assert 0 <= oy and oy + TILE_SPAN <= 1.0001
    assert {(round(x, 2), round(y, 2)) for x, y in TILE_ORIGINS} == {
        (0.0, 0.0), (0.45, 0.0), (0.0, 0.45), (0.45, 0.45)
    }


def test_box_inside_a_quarter_maps_back_to_the_page():
    from app.providers.openai_vision import OpenAIExtractionProvider
    from app.providers.reading import FieldReading

    judge = OpenAIExtractionProvider(api_key="test", model="gpt-4o-mini", split=4)
    # רבע 4 (ימני-תחתון) מתחיל ב-0.45; הנקודה 0.5 בתוכו = 0.45 + 0.5*0.55
    region = judge._region_of(FieldReading(values=["x"], legibility="clear", tile=4, box=[500, 500, 600, 600]))
    assert region == pytest.approx((0.725, 0.725, 0.78, 0.78))


def test_quarters_are_sent_as_four_images(tmp_path):
    from app.providers.openai_vision import page_images

    assert len(page_images(png(tmp_path / "p.png"), 4)) == 4
    assert len(page_images(png(tmp_path / "q.png"), 1)) == 1


def test_gemini_box_order_is_converted():
    from app.providers.gemini_vision import _region_of
    from app.providers.reading import FieldReading

    # Gemini: [ymin, xmin, ymax, xmax]
    region = _region_of(FieldReading(values=["x"], legibility="clear", box=[100, 200, 150, 400]))
    assert region == pytest.approx((0.2, 0.1, 0.4, 0.15))


# --------------------------------------------------------------- the API


@pytest.fixture()
def api(tmp_path, monkeypatch):
    monkeypatch.setattr("app.config.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.store.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.gold.GOLD_DIR", tmp_path / "gold")
    from app.api import app

    batch = Batch(batch_id="bat_api")
    source = SourceFile(index=1, filename="form.png", kind="png", sha256="f" * 64)
    batch.sources.append(source)
    page = pipeline.page_image_path("bat_api", 1, 1)
    page.parent.mkdir(parents=True, exist_ok=True)
    png(page, size=(1000, 1400))
    record = record_from(
        reading("a", regions={"national_id": (0.4, 0.4, 0.6, 0.45)}, national_id=["001122344"],
                commitment_numbers=["730578626"]),
        reading("b", national_id=["00112?344"], commitment_numbers=["730578626"]),
    )
    record.batch_id, record.source_index, record.page_number = "bat_api", 1, 1
    batch.records.append(record)
    batch.status = "completed"
    store.put(batch)
    yield TestClient(app), record
    store.delete("bat_api")


def test_api_serves_candidates_and_crop_box(api):
    client, record = api
    data = client.get("/api/batches/bat_api").json()
    field = data["records"][0]["fields"]["national_id"]
    assert field["candidates"] == [["001122344"], ["00112?344"]]
    assert field["crop_box"] is not None


def test_api_crop_endpoint(api):
    client, _ = api
    res = client.get("/api/batches/bat_api/sources/1/pages/1/crops/national_id")
    assert res.status_code == 200 and res.headers["content-type"] == "image/png"
    assert res.headers["cache-control"] == "no-store"


def test_api_no_crop_without_a_reliable_region(api):
    client, _ = api
    res = client.get("/api/batches/bat_api/sources/1/pages/1/crops/first_name")
    assert res.status_code == 404, "לא ממציאים תקריב"


def test_choosing_an_option_resolves_the_field(api):
    client, record = api
    res = client.patch(f"/api/records/{record.record_id}",
                       json={"field": "national_id", "values": ["001122344"]})
    field = res.json()["records"][0]["fields"]["national_id"]
    assert field["values"] == ["001122344"]
    assert field["state"] == STATE_APPROVED
    assert field["candidates"] == []


def test_verify_is_refused_while_a_choice_is_open(api):
    client, record = api
    res = client.post(f"/api/records/{record.record_id}/verify")
    assert res.status_code == 409
    assert "תעודת זהות" in res.json()["detail"]


def test_verify_saves_an_answer_key_entry(api, tmp_path):
    client, record = api
    client.patch(f"/api/records/{record.record_id}", json={"field": "national_id", "values": ["001122344"]})
    res = client.post(f"/api/records/{record.record_id}/verify")
    assert res.status_code == 200
    assert res.json()["records"][0]["verified"] is True

    forms = gold.load_all(tmp_path / "gold")
    assert len(forms) == 1
    assert forms[0].truth["national_id"] == ["001122344"]
    assert forms[0].page_png.is_file()


def test_holdout_is_deterministic_and_about_a_third():
    picks = [gold.is_holdout(f"{i:064x}", 1) for i in range(300)]
    assert picks == [gold.is_holdout(f"{i:064x}", 1) for i in range(300)]
    assert 70 < sum(picks) < 130


# ----------------------------------------------------------------- Excel


def test_disputed_field_is_empty_in_excel_with_options_in_reasons(tmp_path):
    batch = Batch(batch_id="x", sources=[SOURCE])
    batch.records.append(record_from(
        reading("a", national_id=["001122344"]), reading("b", national_id=["00112?344"])
    ))
    path = tmp_path / "out.xlsx"
    path.write_bytes(workbook_bytes(batch))
    ws = load_workbook(path).active
    headers = [c.value for c in ws[1]]
    assert ws.cell(2, headers.index("תעודת זהות") + 1).value in (None, "")
    reasons = ws.cell(2, headers.index("סיבות") + 1).value
    assert "001122344" in reasons and "00112?344" in reasons


# ------------------------------------------------------------- evaluate


def test_evaluation_counts_a_silent_error(tmp_path):
    from app.evaluate import Config, Score, score_form

    form = gold.GoldForm("k", png(tmp_path / "p.png"), "f.pdf", 1, "", {"commitment_numbers": ["730578626"]}, False)
    wrong = Judge("a", reading("a", commitment_numbers=["87250"]))
    wrong2 = Judge("b", reading("b", commitment_numbers=["87250"]))
    cfg = Config("X", "test", lambda: [wrong, wrong2], 0)
    score = Score()
    score_form(cfg, form, score, [])
    assert score.silent_errors == 1, "שני השופטים הסכימו על ערך שגוי — זה בדיוק מה שנמדד"
    assert score.errors[0]["field"] == "commitment_numbers"


# --------------------------------------------------------- Gemini quotas


class _GeminiError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code
        self.message = message


FREE_TIER = (
    "You exceeded your current quota, please check your plan and billing details. "
    "* Quota exceeded for metric: generativelanguage.googleapis.com/generate_content_free_tier_requests, "
    "limit: 20, model: gemini-3.8-flash Please retry in 4h11m2.9s."
)


def test_free_tier_daily_quota_is_named_and_not_retried():
    from app.providers.gemini_vision import _classify, _is_transient_rate_limit

    exc = _GeminiError(429, FREE_TIER)
    assert not _is_transient_rate_limit(exc), "לא מנסים שוב מכסה שמתאפסת בעוד 4 שעות"
    message, fatal = _classify(exc)
    assert fatal
    assert "20 בקשות ביום" in message and "חיוב" in message
    assert "קרדיט" not in message, "זו לא בעיית קרדיט — ההודעה הקודמת הטעתה"


def test_short_rate_limit_is_still_retried():
    from app.providers.gemini_vision import _is_transient_rate_limit

    exc = _GeminiError(429, "Resource has been exhausted. Please retry in 12s.")
    assert _is_transient_rate_limit(exc)


def test_overload_503_is_retried():
    from app.providers.gemini_vision import _is_transient_rate_limit

    assert _is_transient_rate_limit(_GeminiError(503, "This model is currently experiencing high demand."))


# ---------------------------------------------- zoom location: trusted judge only


def test_gpt_alone_gives_no_zoom_location():
    gpt = reading("gpt", regions={"commitment_numbers": (0.08, 0.44, 0.87, 0.58)},
                  commitment_numbers=["2646"])
    gpt.regions_reliable = False
    r = record_from(gpt)
    assert r.fields["commitment_numbers"].region is None, "בלי מיקום אמין — טופס מלא, לא תקריב שגוי"


def test_gemini_location_is_used_even_when_gpt_is_listed_first():
    gpt = reading("gpt", regions={"commitment_numbers": (0.08, 0.44, 0.87, 0.58)},
                  commitment_numbers=["2646"])
    gpt.regions_reliable = False
    gemini = reading("gemini", regions={"commitment_numbers": (0.18, 0.37, 0.30, 0.40)},
                     commitment_numbers=["2646"])
    r = record_from(gpt, gemini)
    assert r.fields["commitment_numbers"].region == (0.18, 0.37, 0.30, 0.40)


def test_no_fallback_to_gpt_when_gemini_has_no_location_for_the_field():
    gpt = reading("gpt", regions={"sending_entity": (0.74, 0.04, 0.91, 0.13)},
                  sending_entity=["שמיר"])
    gpt.regions_reliable = False
    gemini = reading("gemini", sending_entity=["שמיר"])  # Gemini לא נתן תיבה לשדה הזה
    r = record_from(gemini, gpt)
    assert r.fields["sending_entity"].region is None


def test_openai_judge_marks_its_locations_unreliable(monkeypatch):
    from types import SimpleNamespace

    from app.providers.openai_vision import OpenAIExtractionProvider
    from app.providers.reading import FieldReading, PageReading

    judge = OpenAIExtractionProvider(api_key="test", model="gpt-5.4-mini")
    empty = FieldReading(values=[], legibility="absent")
    parsed = PageReading(template="x", **{
        k: empty for k in PageReading.model_fields if k not in ("template", "issue_date", "birth_date", "note")
    })
    completion = SimpleNamespace(
        choices=[SimpleNamespace(message=SimpleNamespace(parsed=parsed, refusal=None, content=""))],
        usage=None,
    )
    assert judge._to_domain(completion).regions_reliable is False


# ------------------------------------------------- editing the prompt takes effect


def test_edited_prompt_is_not_served_from_the_cache(tmp_path):
    pipeline._READ_CACHE.clear()
    source = SourceFile(index=1, filename="f.pdf", sha256="same-file")
    judge = Judge("a", reading("a", commitment_numbers=["1"]))
    image = png(tmp_path / "p.png")

    judge.prompt_version = "extract_v2-aaaa1111"
    read_form(Panel([judge]), image, "b", source, 1)
    read_form(Panel([judge]), image, "b", source, 1)
    assert judge.calls == 1, "אותו פרומפט — מהמטמון"

    judge.prompt_version = "extract_v2-bbbb2222"  # נועה ערכה את הפרומפט
    read_form(Panel([judge]), image, "b", source, 1)
    assert judge.calls == 2, "פרומפט חדש — קריאה חדשה, לא תשובה ישנה"
    pipeline._READ_CACHE.clear()


def test_prompt_version_follows_the_text():
    from app.providers.reading import prompt_version

    assert prompt_version("א") == prompt_version("א")
    assert prompt_version("א") != prompt_version("ב")


# ------------------------------------------------- model comparison per field


def test_each_field_records_what_every_judge_read_when_they_agree():
    r = record_from(
        reading("gemini-3.8-flash", central_numbers=["939376"]),
        reading("gpt-5.4-mini", central_numbers=["939376"]),
    )
    assert r.fields["central_numbers"].readings == [
        {"model": "gemini-3.8-flash", "values": ["939376"], "legibility": "clear"},
        {"model": "gpt-5.4-mini", "values": ["939376"], "legibility": "clear"},
    ]


def test_readings_are_kept_when_judges_disagree_and_are_normalized():
    r = record_from(
        reading("gemini-3.8-flash", sample_received_date=["18/06/2026"]),
        reading("gpt-5.4-mini", legibility={"sample_received_date": "partial"},
                sample_received_date=["22-08-26"]),
    )
    readings = r.fields["sample_received_date"].readings
    assert [x["values"] for x in readings] == [["18/06/2026"], ["22/08/2026"]]
    assert [x["legibility"] for x in readings] == ["clear", "partial"]


def test_single_judge_has_one_reading():
    r = record_from(reading("gpt-5.4-mini", first_name=["שמו?ל"]))
    assert len(r.fields["first_name"].readings) == 1


def test_partial_name_from_both_judges_is_visible_but_orange():
    r = record_from(
        reading("a", legibility={"first_name": "partial"}, first_name=["שמו?ל"]),
        reading("b", legibility={"first_name": "partial"}, first_name=["שמו?ל"]),
    )
    field = r.fields["first_name"]
    assert field.values == ["שמו?ל"], "המזכירה רואה את מה שנקרא"
    assert field.state == STATE_REVIEW, "אבל זה לא מאושר — כתום"


def test_api_serves_readings(api):
    client, _ = api
    field = client.get("/api/batches/bat_api").json()["records"][0]["fields"]["national_id"]
    assert [x["values"] for x in field["readings"]] == [["001122344"], ["00112?344"]]


# ----------------------------------------------------------- Claude as judge B


def _claude_response(parsed=None, stop="end_turn", model="claude-sonnet-5-5", usage=None):
    from types import SimpleNamespace

    return SimpleNamespace(parsed_output=parsed, stop_reason=stop, model=model, content=[], usage=usage)


def _page_reading(**fields):
    from app.providers.reading import FieldReading, PageReading

    empty = FieldReading(values=[], legibility="absent")
    reading = PageReading(template="kupah_form", **{
        k: empty for k in PageReading.model_fields if k not in ("template", "issue_date", "birth_date", "note")
    })
    for key, value in fields.items():
        setattr(reading, key, value)
    return reading


def test_judge_b_is_chosen_by_model_name():
    from app.pipeline import judge_b_class
    from app.providers.claude_vision import ClaudeExtractionProvider
    from app.providers.openai_vision import OpenAIExtractionProvider

    assert judge_b_class("claude-sonnet-5-5") is ClaudeExtractionProvider
    assert judge_b_class("gpt-5.4-mini") is OpenAIExtractionProvider


def test_claude_answer_becomes_a_reading_with_untrusted_locations():
    from types import SimpleNamespace

    from app.providers.claude_vision import ClaudeExtractionProvider
    from app.providers.reading import FieldReading

    judge = ClaudeExtractionProvider(api_key="test", model="claude-sonnet-5-5")
    parsed = _page_reading(
        national_id=FieldReading(values=["0011223344"], legibility="clear", tile=1, box=[100, 300, 400, 340]),
    )
    usage = SimpleNamespace(input_tokens=5000, output_tokens=800,
                            cache_read_input_tokens=6000, cache_creation_input_tokens=0)
    raw = judge._to_domain(_claude_response(parsed, usage=usage))

    assert raw.values["national_id"] == ["0011223344"]
    assert raw.regions_reliable is False, "ההגדלה ממשיכה לקחת מיקום רק מ-Gemini"
    assert raw.usage.cost_usd == pytest.approx((5000 * 2 + 6000 * 0.2 + 800 * 10) / 1_000_000)


def test_claude_refusal_does_not_stop_the_batch():
    from app.providers.claude_vision import ClaudeExtractionProvider

    judge = ClaudeExtractionProvider(api_key="test")
    with pytest.raises(ExtractionError) as caught:
        judge._to_domain(_claude_response(stop="refusal"))
    assert caught.value.fatal is False


def test_claude_fallback_model_is_named_and_priced():
    from app.providers.claude_vision import ClaudeExtractionProvider

    judge = ClaudeExtractionProvider(api_key="test", model="claude-sonnet-5-5")
    raw = judge._to_domain(_claude_response(_page_reading(), model="claude-opus-5-5"))
    assert raw.model == "claude-opus-5-5", "אם עבר למודל גיבוי — רואים מי ענה בפועל"


def _api_error(cls, message, status):
    import httpx2

    request = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
    return cls(message, response=httpx2.Response(status, request=request), body=None)


def test_claude_account_problems_are_fatal_and_overload_is_not():
    import anthropic

    from app.providers.claude_vision import _classify

    assert _classify(_api_error(anthropic.AuthenticationError, "invalid x-api-key", 401))[1] is True
    assert _classify(_api_error(anthropic.BadRequestError, "Your credit balance is too low", 400))[1] is True
    assert _classify(_api_error(anthropic.NotFoundError, "model not found", 404))[1] is True
    assert _classify(_api_error(anthropic.RateLimitError, "rate limited", 429))[1] is False
    assert _classify(_api_error(anthropic.InternalServerError, "overloaded", 529))[1] is False


def test_page_sent_to_claude_fits_its_limit(tmp_path):
    import base64
    import io

    from app.providers.claude_vision import MAX_LONG_SIDE, page_png_b64

    page = png(tmp_path / "p.png", size=(2481, 3509))
    data = base64.b64decode(page_png_b64(page))
    with Image.open(io.BytesIO(data)) as sent:
        assert max(sent.size) == MAX_LONG_SIDE
        assert sent.size[0] / sent.size[1] == pytest.approx(2481 / 3509, rel=0.01)


# ------------------------------------------------------------- "bold ink" version

INK = (74, 187, 237)  # נמדד על טופס מאוחדת
TURQUOISE = (168, 247, 247)
YELLOW = (245, 253, 171)
PAPER = (251, 253, 251)
PRINT = (34, 31, 36)


def _swatches(path: Path) -> Path:
    image = Image.new("RGB", (500, 100), PAPER)
    for i, colour in enumerate((INK, TURQUOISE, YELLOW, PAPER, PRINT)):
        image.paste(colour, (i * 100, 0, i * 100 + 100, 100))
    image.save(path)
    return path


def test_ink_version_removes_highlighters_and_darkens_ink(tmp_path):
    from app.pages import ink_enhanced

    out = ink_enhanced(_swatches(tmp_path / "page-001.png"))
    with Image.open(out) as ink:
        values = [ink.getpixel((i * 100 + 50, 50)) for i in range(5)]
    blue, turquoise, yellow, paper, printed = values
    assert blue < 70, "עט כחול נעשה כהה"
    assert turquoise == 255 and yellow == 255, "המרקרים נעלמים"
    assert paper == 255
    assert printed == 0, "דפוס שחור נשאר שחור"


def test_ink_version_keeps_the_page_size_and_is_made_once(tmp_path):
    from app.pages import ink_enhanced

    page = png(tmp_path / "page-001.png", size=(1234, 1700))
    out = ink_enhanced(page)
    assert out.name == "page-001.ink.png"
    with Image.open(out) as ink:
        assert ink.size == (1234, 1700), "אותו גודל — המיקומים לא זזים"
    stamp = out.stat().st_mtime_ns
    ink_enhanced(page)
    assert out.stat().st_mtime_ns == stamp


def test_prompt_mentions_the_second_image_only_when_it_is_sent():
    from app.providers.reading import load_prompt, prompt_version

    plain, with_ink = load_prompt("x"), load_prompt("x", with_ink=True)
    assert "שתי תמונות של אותו עמוד" in with_ink
    assert "שתי תמונות של אותו עמוד" not in plain
    assert prompt_version(plain) != prompt_version(with_ink), "המטמון לא יחזיר תשובה ישנה"


def _settings(**changes):
    from app.config import SETTINGS

    return replace(SETTINGS, **changes)


@pytest.mark.parametrize("enabled, expected", [(True, 2), (False, 1)])
def test_gemini_sends_the_ink_version(tmp_path, monkeypatch, enabled, expected):
    from types import SimpleNamespace

    from app.providers import gemini_vision as gv

    sent = {}

    class _Models:
        def generate_content(self, model, contents, config):
            sent["images"] = [c for c in contents if not isinstance(c, str)]
            return SimpleNamespace(parsed=_page_reading(), text="", usage_metadata=None)

    monkeypatch.setattr(gv, "SETTINGS", _settings(enhance_ink=enabled))
    monkeypatch.setattr(gv.genai, "Client", lambda api_key: SimpleNamespace(models=_Models()))
    gv.GeminiExtractionProvider(api_key="test").extract(png(tmp_path / "page-001.png"))
    assert len(sent["images"]) == expected


@pytest.mark.parametrize("enabled, expected", [(True, 2), (False, 1)])
def test_claude_sends_the_ink_version_first_original_then_ink(tmp_path, monkeypatch, enabled, expected):
    from types import SimpleNamespace

    from app.providers import claude_vision as cv

    sent = {}

    def parse(**kwargs):
        sent["content"] = kwargs["messages"][0]["content"]
        return _claude_response(_page_reading())

    monkeypatch.setattr(cv, "SETTINGS", _settings(enhance_ink=enabled))
    judge = cv.ClaudeExtractionProvider(api_key="test")
    judge._client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(parse=parse)))
    page = _swatches(tmp_path / "page-001.png")
    judge.extract(page)

    images = [c for c in sent["content"] if c["type"] == "image"]
    assert len(images) == expected
    assert sent["content"][-1]["type"] == "text"
    if enabled:
        from app.providers.claude_vision import page_png_b64

        assert images[0]["source"]["data"] == page_png_b64(page), "המקורית ראשונה"


def test_gpt_quarters_do_not_get_an_ink_copy(monkeypatch):
    from app.providers import openai_vision as ov

    monkeypatch.setattr(ov, "SETTINGS", _settings(enhance_ink=True))
    assert ov.OpenAIExtractionProvider(api_key="t", model="gpt-4o-mini", split=4).enhance_ink is False
    assert ov.OpenAIExtractionProvider(api_key="t", model="gpt-5.4-mini", split=1).enhance_ink is True


def test_api_crop_in_ink_view(api):
    client, _ = api
    # עמוד עם מרקר טורקיז — בגרסה המודגשת הוא אמור להיעלם
    Image.new("RGB", (1000, 1400), TURQUOISE).save(pipeline.page_image_path("bat_api", 1, 1))
    plain = client.get("/api/batches/bat_api/sources/1/pages/1/crops/national_id")
    ink = client.get("/api/batches/bat_api/sources/1/pages/1/crops/national_id?ink=1")
    assert plain.status_code == ink.status_code == 200
    assert ink.headers["content-type"] == "image/png"
    assert ink.content != plain.content, "גרסה נפרדת, לא אותה תמונה"
