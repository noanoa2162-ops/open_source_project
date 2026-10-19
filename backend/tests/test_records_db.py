"""Records live in the database: they survive a restart, can be searched and
filtered, belong to one browser session, and are exported as safe text."""

from __future__ import annotations

import io
import json
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import api, db, pipeline, store
from app.models import Batch, FormRecord, SourceFile

SOURCE = SourceFile(index=1, filename="part1.pdf", kind="pdf", sha256="b" * 64)


def _record(batch_id: str, page: int = 1, **values) -> FormRecord:
    record = FormRecord.empty(batch_id, SOURCE, page)
    base = {
        "first_name": ["ישראל"], "last_name": ["ישראלי"], "national_id": ["011223344"],
        "central_numbers": ["939376"], "sample_received_date": ["18/06/2026"],
        "sending_entity": ["מאוחדת"], "commitment_numbers": ["0730578626"],
    }
    base.update(values)
    for key, vals in base.items():
        record.fields[key].values = list(vals)
    record.fields["national_id"].region = (0.1, 0.2, 0.3, 0.25)
    return record


def _batch(batch_id: str, session: str, *records: FormRecord, status: str = "completed") -> Batch:
    batch = Batch(batch_id=batch_id, session_id=session, status=status)
    batch.sources.append(SOURCE)
    batch.records.extend(records)
    return batch


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr("app.config.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.store.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.pipeline.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.gold.GOLD_DIR", tmp_path / "gold")
    created = []
    yield created
    for batch_id in created:
        store.delete(batch_id)


def _client(session: str) -> TestClient:
    client = TestClient(api.app)
    client.cookies.set(api.SESSION_COOKIE, session)
    return client


SID_A = "a" * 43
SID_B = "b" * 43


# ------------------------------------------------------------ המסד עצמו
def test_a_batch_survives_a_restart_exactly():
    record = _record("bat_r1")
    record.fields["first_name"].state = "needs_review"
    record.fields["first_name"].reasons = ["illegible"]
    record.revision = 3
    db.save_batch(_batch("bat_r1", SID_A, record))

    loaded = {b.batch_id: b for b in db.load_batches()}["bat_r1"]
    back = loaded.records[0]
    assert loaded.session_id == SID_A and loaded.sources[0].filename == "part1.pdf"
    assert back.fields["national_id"].values == ["011223344"], "אפס מוביל נשמר"
    assert back.fields["national_id"].region == (0.1, 0.2, 0.3, 0.25), "מיקום חוזר כ-tuple"
    assert back.fields["first_name"].reasons == ["illegible"] and back.revision == 3


def test_saving_again_never_duplicates_and_drops_removed_records():
    one, two = _record("bat_r2", 1), _record("bat_r2", 2)
    batch = _batch("bat_r2", SID_A, one, two)
    db.save_batch(batch)
    db.save_batch(batch)
    assert len(db.search_records(SID_A)) == 2
    batch.records.remove(two)
    db.save_batch(batch)
    assert [r["record_id"] for r in db.search_records(SID_A)] == [one.record_id]


def test_search_and_filters():
    a = _record("bat_r3", 1)
    b = _record("bat_r3", 2, sending_entity=["מכבי"], sample_received_date=["01/06/2026"],
                central_numbers=["937309", "937310"])
    b.verified = True
    c = _record("bat_r3", 3, first_name=["דוד"])
    c.fields["first_name"].state = "needs_review"
    db.save_batch(_batch("bat_r3", SID_A, a, b, c))

    ids = lambda rows: {r["record_id"] for r in rows}  # noqa: E731
    assert ids(db.search_records(SID_A, q="937310")) == {b.record_id}, "חיפוש במספר מרכזי מתוך רשימה"
    assert ids(db.search_records(SID_A, q="דוד")) == {c.record_id}
    assert ids(db.search_records(SID_A, status="approved")) == {b.record_id}
    assert ids(db.search_records(SID_A, status="review")) == {c.record_id}
    assert ids(db.search_records(SID_A, sender="מכבי")) == {b.record_id}
    assert ids(db.search_records(SID_A, date_from="2026-06-10")) == {a.record_id, c.record_id}
    assert ids(db.search_records(SID_A, date_to="2026-06-05")) == {b.record_id}
    counts = db.record_counts(SID_A)
    assert counts == {"ready": 1, "approved": 1, "review": 1, "all": 3}
    assert db.search_records(SID_B) == [], "session אחר לא רואה כלום"


def test_duplicate_numbers_are_found_within_the_session_only():
    first = _record("bat_d1", 1)
    db.save_batch(_batch("bat_d1", SID_A, first))
    db.save_batch(_batch("bat_d9", SID_B, _record("bat_d9", 1)))
    again = _record("bat_d2", 1)
    db.save_batch(_batch("bat_d2", SID_A, again))
    dups = db.find_duplicates(again, SID_A)
    assert [d["record_id"] for d in dups] == [first.record_id]
    assert set(dups[0]["matches"]) == {"939376", "0730578626"}
    clean = _record("bat_d3", 1, central_numbers=["111111"], commitment_numbers=["22222"])
    assert db.find_duplicates(clean, SID_A) == []


# ------------------------------------------------------------ ה-API
def test_each_session_sees_only_its_own_uploads(isolated):
    record = _record("bat_s1")
    batch = _batch("bat_s1", SID_A, record)
    isolated.append("bat_s1")
    store.put(batch)
    store.persist(batch)
    page = pipeline.page_image_path("bat_s1", 1, 1)
    page.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (60, 80), "white").save(page)

    owner, stranger = _client(SID_A), _client(SID_B)
    assert owner.get("/api/batches/bat_s1").status_code == 200
    assert stranger.get("/api/batches/bat_s1").status_code == 404
    assert stranger.get("/api/batches/bat_s1/sources/1/pages/1").status_code == 404
    assert stranger.patch(f"/api/records/{record.record_id}", json={"field": "first_name", "approve": True}).status_code == 404
    assert stranger.post(f"/api/records/{record.record_id}/verify").status_code == 404
    assert stranger.delete("/api/batches/bat_s1").status_code == 404
    assert stranger.get("/api/records").json()["records"] == []
    assert [r["record_id"] for r in owner.get("/api/records").json()["records"]] == [record.record_id]
    assert [b["batch_id"] for b in owner.get("/api/batches").json()] == ["bat_s1"]


def test_a_new_browser_gets_a_secret_session_cookie():
    response = TestClient(api.app).get("/api/config")
    cookie = response.headers["set-cookie"]
    assert api.SESSION_COOKIE in cookie and "HttpOnly" in cookie and "samesite=lax" in cookie.lower()


def test_a_fix_is_saved_to_the_database_and_a_stale_screen_gets_409(isolated):
    record = _record("bat_s2")
    batch = _batch("bat_s2", SID_A, record)
    isolated.append("bat_s2")
    store.put(batch)
    client = _client(SID_A)

    body = {"field": "last_name", "values": ["כהן"], "revision": 0}
    assert client.patch(f"/api/records/{record.record_id}", json=body).status_code == 200
    stale = client.patch(f"/api/records/{record.record_id}", json={**body, "values": ["לוי"]})
    assert stale.status_code == 409, "גרסה ישנה — נדחה, לא דורס"

    saved = db.search_records(SID_A, q="כהן")
    assert [r["record_id"] for r in saved] == [record.record_id] and saved[0]["revision"] == 1
    reloaded = {b.batch_id: b for b in db.load_batches()}["bat_s2"]
    assert reloaded.records[0].fields["last_name"].values == ["כהן"], "התיקון שרד הפעלה מחדש"


def test_restart_reloads_batches_and_resumes_unfinished_ones(isolated, monkeypatch):
    done = _batch("bat_s3", SID_A, _record("bat_s3"))
    running = _batch("bat_s4", SID_A, status="processing")
    isolated.extend(["bat_s3", "bat_s4"])
    db.save_batch(done)
    db.save_batch(running)
    resumed = []
    monkeypatch.setattr(pipeline, "start", lambda batch_id: resumed.append(batch_id))
    for batch in store.restore():
        pipeline.resume(batch)
    assert store.get("bat_s3") is not None and resumed == ["bat_s4"]
    assert store.get("bat_s4").status == "queued" and store.get("bat_s4").warnings


def test_records_export_keeps_text_and_neutralises_formulas(isolated):
    record = _record("bat_s5", commitment_numbers=["0012", "=HYPERLINK(1)"])
    batch = _batch("bat_s5", SID_A, record)
    isolated.append("bat_s5")
    store.put(batch)
    store.persist(batch)
    client = _client(SID_A)

    xlsx = client.get("/api/records/export?format=xlsx")
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    assert len(rows) == 2, "שורה לכל מספר התחייבות"
    assert rows[0][6] == "0012" and rows[1][6] == "'=HYPERLINK(1)"
    assert rows[0][2] == "011223344"

    csv_text = client.get("/api/records/export?format=csv").content.decode("utf-8-sig")
    assert "'=HYPERLINK(1)" in csv_text and "011223344" in csv_text

    data = json.loads(client.get("/api/records/export?format=json").content)
    assert data[0]["commitment_numbers"] == ["0012", "=HYPERLINK(1)"], "JSON שומר את הערך המקורי"


def test_access_code_guards_everything_but_login(monkeypatch):
    monkeypatch.setattr(api, "SETTINGS", replace(api.SETTINGS, access_code="s3cret"))
    client = TestClient(api.app)
    assert client.get("/api/records").status_code == 401
    assert client.get("/", follow_redirects=False).status_code == 303
    assert client.get("/healthz").status_code == 200
    assert client.post("/login", data={"code": "wrong"}).status_code == 401
    assert client.post("/login", data={"code": "s3cret"}, follow_redirects=False).status_code == 303
    assert client.get("/api/records").status_code == 200


def test_changes_from_another_site_are_refused():
    client = TestClient(api.app)
    response = client.patch("/api/records/rec_x", json={"field": "first_name", "approve": True},
                            headers={"Origin": "https://evil.example"})
    assert response.status_code == 403
