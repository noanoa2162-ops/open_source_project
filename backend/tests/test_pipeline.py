"""Checks for the rules that would quietly corrupt a form if they broke.

Nothing here calls OpenAI. A stub provider stands in for the model so the
pipeline, the validators and the Excel writer can be exercised offline.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.config import SETTINGS
from app.excel import neutralize, workbook_bytes
from app.fields import FIELDS
from app.files import KIND_JPEG, KIND_PDF, KIND_PNG, sniff_kind
from app.models import STATE_OK, STATE_REVIEW, Batch, FormRecord, SourceFile
from app.providers.base import ExtractionError, RawExtraction
from app.validators import validate_record


class StubProvider:
    """Stands in for a judge. Returns whatever the test hands it."""

    name = "stub"
    model = "stub"

    def __init__(self, values: dict, note: str = "") -> None:
        self._values = values
        self._note = note

    def extract(self, image_path: Path) -> RawExtraction:
        return RawExtraction(values=self._values, note=self._note, model="stub")


SOURCE = SourceFile(index=1, filename="sample.pdf", kind="pdf")


def _record(**values) -> FormRecord:
    record = FormRecord.empty("b1", SOURCE, 1)
    for key, vals in values.items():
        record.fields[key].values = vals
    return record


# ---------------------------------------------------------------- uploads


def test_file_kind_comes_from_bytes_not_extension():
    assert sniff_kind(b"%PDF-1.7 ...") == KIND_PDF
    assert sniff_kind(b"\xff\xd8\xff\xe0 jfif") == KIND_JPEG
    assert sniff_kind(b"\x89PNG\r\n\x1a\n") == KIND_PNG
    # a text file renamed to .pdf must still be rejected
    assert sniff_kind(b"hello world") is None


# ---------------------------------------------------------------- validators


def test_empty_field_is_flagged_not_hidden():
    record = _record()
    validate_record(record)
    assert all(f.state == STATE_REVIEW for f in record.fields.values())
    assert record.fields["national_id"].reasons == ["missing"]


def test_valid_id_and_date_pass_clean():
    record = _record(national_id=["200000008"], sample_received_date=["18/06/2026"])
    validate_record(record)
    assert record.fields["national_id"].state == STATE_OK
    assert record.fields["sample_received_date"].state == STATE_OK


def test_leading_zeros_are_never_stripped():
    record = _record(national_id=["012345678"])
    validate_record(record)
    assert record.fields["national_id"].values == ["012345678"]


def test_bad_id_is_flagged_but_value_kept():
    record = _record(national_id=["12345"])
    validate_record(record)
    field = record.fields["national_id"]
    assert field.state == STATE_REVIEW
    assert "id_length" in field.reasons
    assert field.values == ["12345"], "ערך לא תקין נשמר כפי שנקרא ולא נמחק"


def test_wrong_date_format_is_flagged():
    record = _record(sample_received_date=["3.10.26"])
    validate_record(record)
    assert "bad_date" in record.fields["sample_received_date"].reasons


def test_impossible_date_is_flagged():
    record = _record(sample_received_date=["32/13/2026"])
    validate_record(record)
    assert "impossible_date" in record.fields["sample_received_date"].reasons


def test_multiple_values_are_kept_and_flagged():
    record = _record(commitment_numbers=["730578626", "730578627"])
    validate_record(record)
    field = record.fields["commitment_numbers"]
    assert field.values == ["730578626", "730578627"]
    assert "multiple_values" in field.reasons


def test_model_note_marks_the_field_it_names():
    record = _record(first_name=["שרה"], national_id=["200000008"])
    validate_record(record, "השם הפרטי מטושטש ולא ברור")
    assert "model_uncertain" in record.fields["first_name"].reasons
    assert record.fields["national_id"].state == STATE_OK


# ---------------------------------------------------------------- excel


def test_formula_values_are_neutralised():
    for dangerous in ("=SUM(A1)", "+1+1", "-2", "@cmd"):
        assert neutralize(dangerous).startswith("'")
    assert neutralize("730578626") == "730578626"


def test_export_keeps_everything_as_text(tmp_path):
    batch = Batch(batch_id="b1", sources=[SOURCE])
    record = _record(
        national_id=["012345678"],
        central_numbers=["939376"],
        commitment_numbers=["730578626", "730578627"],
        sample_received_date=["18/06/2026"],
        first_name=["=SUM(A1)"],
    )
    validate_record(record)
    batch.records.append(record)

    path = tmp_path / "out.xlsx"
    path.write_bytes(workbook_bytes(batch))
    ws = load_workbook(path).active

    headers = [ws.cell(1, i + 1).value for i in range(len(FIELDS))]
    assert headers == [f.label_he for f in FIELDS]

    row = {ws.cell(1, i).value: ws.cell(2, i) for i in range(1, len(FIELDS) + 1)}
    assert row["תעודת זהות"].value == "012345678"
    assert row["תעודת זהות"].number_format == "@"
    assert row["תאריך קבלת הדגימה"].value == "18/06/2026"
    # דרישת הפרויקט: שורה לכל מספר התחייבות
    assert row["מספר התחייבות"].value == "730578626"
    assert ws.cell(3, list(row).index("מספר התחייבות") + 1).value == "730578627"
    assert row["שם פרטי"].value.startswith("'"), "נוסחה חייבת להינטרל"


# ---------------------------------------------------------------- provider


def test_fatal_error_carries_its_flag():
    assert ExtractionError("no credit", fatal=True).fatal is True
    assert ExtractionError("one bad page").fatal is False


def test_stub_provider_fills_the_record(tmp_path):
    from app.pipeline import Reader, read_form

    provider = StubProvider(
        {
            "first_name": ["שרה"],
            "national_id": ["200000008"],
            "commitment_numbers": ["730578626"],
            "central_numbers": [],
            "last_name": ["לוי"],
            "sample_received_date": ["18/06/2026"],
            "sending_entity": ["מאוחדת"],
        }
    )
    record = read_form(Reader(provider), _png(tmp_path / "page-001.png"), "b1", SOURCE, 1)

    assert record.error == ""
    assert record.fields["first_name"].values == ["שרה"]
    assert record.fields["sending_entity"].values == ["מאוחדת"]
    # שדה שהמודל לא מצא נשאר ריק ומסומן, לא מומצא
    assert record.fields["central_numbers"].values == []
    assert record.fields["central_numbers"].state == STATE_REVIEW


def test_provider_failure_does_not_lose_the_record(tmp_path):
    from app.pipeline import Reader, read_form

    class Broken:
        name = "broken"
        model = "broken"

        def extract(self, image_path):
            raise ExtractionError("אין יתרת קרדיט", fatal=True)

    reader = Reader(Broken())
    record = read_form(reader, _png(tmp_path / "p.png"), "b1", SOURCE, 3)
    assert reader.exhausted, "שגיאת חשבון עוצרת את הקורא"
    assert record.page_number == 3
    assert "קרדיט" in record.error
    assert record.review_count == len(FIELDS)


# ---------------------------------------------------------------- many files

# אלה הבדיקות שמוכיחות את התיקון: אותו מספר טפסים חייב לתת אותו מספר שורות,
# בלי קשר לכמה קבצים הם הגיעו ומאיזה סוג.


def _png(path: Path, size=(40, 56)) -> Path:
    Image.new("RGB", size, "white").save(path, format="PNG")
    return path


def _pdf(path: Path, pages: int) -> Path:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument.new()
    for _ in range(pages):
        doc.new_page(200, 280).close()
    doc.save(path)
    doc.close()
    return path


def _drop(
    tmp_path: Path, monkeypatch, files: list[Path], provider=None, parallel: int | None = None
) -> Batch:
    """Simulate one drop of several files, all the way through the pipeline."""
    from dataclasses import replace

    from app import pipeline, store

    monkeypatch.setattr("app.config.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.store.BATCHES_DIR", tmp_path / "batches")
    chosen = provider or StubProvider({"first_name": ["שרה"]})
    monkeypatch.setattr(pipeline, "build_reader", lambda *_a, **_k: (chosen, ""))
    if parallel is not None:
        monkeypatch.setattr(
            pipeline, "SETTINGS", replace(pipeline.SETTINGS, max_parallel_calls=parallel)
        )

    batch = Batch(batch_id="bat_test")
    for position, path in enumerate(files, start=1):
        batch.sources.append(
            SourceFile(index=position, filename=path.name, kind=path.suffix.lstrip("."))
        )
        (store.source_dir(batch.batch_id, position) / "source").write_bytes(
            path.read_bytes()
        )
    store.put(batch)
    pipeline._process(batch.batch_id)
    return batch


def test_three_images_give_three_rows(tmp_path, monkeypatch):
    files = [_png(tmp_path / f"form{i}.png") for i in range(1, 4)]
    batch = _drop(tmp_path, monkeypatch, files)

    assert batch.status == "completed"
    assert batch.total_pages == 3
    assert len(batch.records) == 3
    assert [r.source_name for r in batch.records] == ["form1.png", "form2.png", "form3.png"]
    assert [r.page_number for r in batch.records] == [1, 1, 1]


def test_one_pdf_of_three_pages_gives_three_rows(tmp_path, monkeypatch):
    batch = _drop(tmp_path, monkeypatch, [_pdf(tmp_path / "many.pdf", 3)])

    assert len(batch.records) == 3
    assert [r.page_number for r in batch.records] == [1, 2, 3]
    assert {r.source_name for r in batch.records} == {"many.pdf"}


def test_image_plus_two_page_pdf_gives_three_rows(tmp_path, monkeypatch):
    files = [_png(tmp_path / "scan.png"), _pdf(tmp_path / "pair.pdf", 2)]
    batch = _drop(tmp_path, monkeypatch, files)

    assert len(batch.records) == 3
    assert [(r.source_name, r.page_number) for r in batch.records] == [
        ("scan.png", 1),
        ("pair.pdf", 1),
        ("pair.pdf", 2),
    ]


def test_all_three_shapes_agree(tmp_path, monkeypatch):
    """The whole point: 3 forms is 3 rows, however they were packaged."""
    a = _drop(tmp_path / "a", monkeypatch, [_png(tmp_path / f"i{i}.png") for i in range(3)])
    b = _drop(tmp_path / "b", monkeypatch, [_pdf(tmp_path / "three.pdf", 3)])
    c = _drop(
        tmp_path / "c", monkeypatch, [_png(tmp_path / "one.png"), _pdf(tmp_path / "two.pdf", 2)]
    )
    assert len(a.records) == len(b.records) == len(c.records) == 3


def test_a_rejected_file_does_not_sink_the_others(tmp_path, monkeypatch):
    from app import pipeline, store

    monkeypatch.setattr("app.config.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.store.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr(
        pipeline, "build_reader", lambda *_a, **_k: (StubProvider({"first_name": ["שרה"]}), "")
    )

    batch = Batch(batch_id="bat_mixed")
    batch.sources.append(SourceFile(index=1, filename="good.png", kind="png"))
    (store.source_dir("bat_mixed", 1) / "source").write_bytes(
        _png(tmp_path / "good.png").read_bytes()
    )
    # קובץ שנדחה כבר בשלב ההעלאה — נושא error ואינו מעובד
    batch.sources.append(
        SourceFile(index=2, filename="notes.txt", error="סוג לא נתמך")
    )
    batch.sources.append(SourceFile(index=3, filename="also.png", kind="png"))
    (store.source_dir("bat_mixed", 3) / "source").write_bytes(
        _png(tmp_path / "also.png").read_bytes()
    )
    store.put(batch)
    pipeline._process("bat_mixed")

    assert len(batch.records) == 2, "שני הקבצים התקינים עובדו"
    assert batch.status == "partial", "הדחייה חייבת להיות גלויה, לא בשקט"
    assert "נדחו" in batch.error
    assert [s.filename for s in batch.rejected_sources] == ["notes.txt"]


def test_page_cap_applies_to_the_whole_drop(tmp_path, monkeypatch):
    from dataclasses import replace

    monkeypatch.setattr(
        "app.pipeline.SETTINGS",
        replace(SETTINGS, page_render_dpi=72, max_pages_per_batch=4),
    )

    files = [_pdf(tmp_path / "a.pdf", 3), _pdf(tmp_path / "b.pdf", 3)]
    batch = _drop(tmp_path, monkeypatch, files)

    assert len(batch.records) == 4, "התקרה חלה על הגרירה כולה, לא על כל קובץ"
    assert batch.skipped_pages == 2, "העמודים שהושמטו חייבים להיספר ולהיות גלויים"


def test_export_names_the_source_file(tmp_path, monkeypatch):
    files = [_png(tmp_path / "alpha.png"), _png(tmp_path / "beta.png")]
    batch = _drop(tmp_path, monkeypatch, files)

    path = tmp_path / "out.xlsx"
    path.write_bytes(workbook_bytes(batch))
    ws = load_workbook(path).active

    headers = [c.value for c in ws[1]]
    assert "קובץ מקור" in headers
    column = headers.index("קובץ מקור") + 1
    assert [ws.cell(r, column).value for r in (2, 3)] == ["alpha.png", "beta.png"]


# ---------------------------------------------------------------- speed


class SlowProvider:
    """Simulates the network wait of a model call. Costs nothing."""

    name = "slow"
    model = "slow"

    def __init__(self, delay_for, fail_for=None, fatal_for=None):
        import threading

        self._delay_for = delay_for
        self._fail_for = fail_for or set()
        self._fatal_for = fatal_for or set()
        self.calls = 0
        self._lock = threading.Lock()

    @staticmethod
    def _source_of(image_path: Path) -> int:
        return int(image_path.parent.parent.name)  # .../sources/<N>/pages/page-001.png

    def extract(self, image_path: Path) -> RawExtraction:
        import time

        n = self._source_of(image_path)
        with self._lock:
            self.calls += 1
        time.sleep(self._delay_for(n))
        if n in self._fatal_for:
            raise ExtractionError("אין יתרת קרדיט", fatal=True)
        if n in self._fail_for:
            raise ExtractionError("עמוד אחד נכשל")
        return RawExtraction(values={"first_name": [f"טופס{n}"]}, model="slow")


def test_pages_are_read_in_parallel(tmp_path, monkeypatch):
    import time

    files = [_png(tmp_path / f"f{i}.png") for i in range(8)]
    started = time.perf_counter()
    batch = _drop(tmp_path, monkeypatch, files, SlowProvider(lambda n: 0.3), parallel=4)
    elapsed = time.perf_counter() - started

    assert len(batch.records) == 8
    # טורי היה לוקח ~2.4 שניות. במקביל של 4: ~0.6. הגבול נדיב כדי שלא יהיה שביר.
    assert elapsed < 1.5, f"נראה שהקריאות רצות בזו אחר זו ({elapsed:.2f}s)"


def test_parallel_of_one_is_still_correct(tmp_path, monkeypatch):
    files = [_png(tmp_path / f"f{i}.png") for i in range(3)]
    batch = _drop(tmp_path, monkeypatch, files, SlowProvider(lambda n: 0.01), parallel=1)
    assert [r.source_name for r in batch.records] == ["f0.png", "f1.png", "f2.png"]


def test_order_is_the_drop_order_not_the_finish_order(tmp_path, monkeypatch):
    files = [_png(tmp_path / f"f{i}.png") for i in range(1, 6)]
    # הקובץ הראשון הכי איטי, האחרון הכי מהיר — הם מסיימים בסדר הפוך
    provider = SlowProvider(lambda n: (6 - n) * 0.08)
    batch = _drop(tmp_path, monkeypatch, files, provider, parallel=5)

    assert [r.source_name for r in batch.records] == [f"f{i}.png" for i in range(1, 6)]
    assert [r.fields["first_name"].values for r in batch.records] == [
        [f"טופס{i}"] for i in range(1, 6)
    ], "כל ערך חייב להישאר צמוד לטופס שלו"


def test_one_failing_page_does_not_stop_the_others(tmp_path, monkeypatch):
    files = [_png(tmp_path / f"f{i}.png") for i in range(1, 5)]
    provider = SlowProvider(lambda n: 0.02, fail_for={2})
    batch = _drop(tmp_path, monkeypatch, files, provider, parallel=4)

    assert len(batch.records) == 4
    assert batch.status == "partial"
    assert [bool(r.error) for r in batch.records] == [False, True, False, False]


def test_fatal_error_stops_new_calls(tmp_path, monkeypatch):
    files = [_png(tmp_path / f"f{i}.png") for i in range(1, 7)]
    provider = SlowProvider(lambda n: 0.02, fatal_for={2})
    # עובד אחד: הסדר דטרמיניסטי, ולכן אפשר לוודא שלא נשלחו קריאות מיותרות
    batch = _drop(tmp_path, monkeypatch, files, provider, parallel=1)

    assert batch.status == "partial", "עמוד 1 נקרא — לא מסתירים אותו מאחורי 'נכשל'"
    assert "קרדיט" in batch.error
    assert provider.calls == 2, "אחרי שגיאת חשבון אסור לשרוף קריאות נוספות"


def test_progress_counts_every_finished_page(tmp_path, monkeypatch):
    files = [_png(tmp_path / f"f{i}.png") for i in range(5)]
    batch = _drop(tmp_path, monkeypatch, files, SlowProvider(lambda n: 0.01), parallel=3)
    assert batch.processed_pages == batch.total_pages == 5


# ---------------------------------------------------------------- provider




def test_page_count_is_known_before_any_page_is_prepared(tmp_path):
    """המסך מראה אחוזים בזמן הזיהוי — לכן הסך הכולל נספר מיד, בלי לרנדר."""
    from app.pages import count_pages

    assert count_pages(_pdf(tmp_path / "three.pdf", 3)) == 3
    image = tmp_path / "one.png"
    Image.new("RGB", (100, 140), "white").save(image)
    assert count_pages(image) == 1
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.7 not really")
    assert count_pages(broken) == 0
